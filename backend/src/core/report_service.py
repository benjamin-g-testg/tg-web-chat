"""Generate downloadable Word and PDF reports from the deterministic analysis."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


REPORTS = (
    ("phase_1", "REPORTE_FASE_1", "Fase 1 - Perfil y evidencia"),
    ("phase_2", "REPORTE_FASE_2", "Fase 2 - Hallazgos e insights"),
    ("phase_3", "REPORTE_FASE_3", "Fase 3 - Respuesta a la consulta"),
    ("conclusion", "REPORTE_CONCLUSION", "Conclusión y recomendaciones"),
    ("executive_summary", "REPORTE_Resumen_Ejecutivo", "Resumen ejecutivo"),
)


def generate_reports(output_dir: Path, reports: dict[str, str], context: dict[str, object]) -> dict[str, dict[str, str]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    artifacts: dict[str, dict[str, str]] = {"word": {}, "pdf": {}}
    for key, filename, title in REPORTS:
        content = str(reports.get(key, "Sin contenido disponible."))
        word_path = output_dir / f"{filename}.docx"
        pdf_path = output_dir / f"{filename}.pdf"
        _write_docx(word_path, title, content, context)
        _write_pdf(pdf_path, title, content, context)
        artifacts["word"][key] = word_path.name
        artifacts["pdf"][key] = pdf_path.name
    return artifacts


def _write_docx(path: Path, title: str, content: str, context: dict[str, object]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    heading = document.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Análisis de calidad Jira", style="Subtitle")
    document.add_heading("Alcance y evidencia", level=2)
    metrics = context.get("metrics", {})
    selected = context.get("selected_issues", [])
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Indicador"
    table.rows[0].cells[1].text = "Resultado"
    for label, value in (
        ("Incidencias en el alcance", metrics.get("total", len(selected))),
        ("Estados reales", metrics.get("status_distribution", {})),
        ("Responsables informados", sum(bool(item.get("assignee")) for item in selected if isinstance(item, dict))),
    ):
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
    document.add_heading("Resultado", level=2)
    for paragraph in _paragraphs(content):
        document.add_paragraph(paragraph)
    document.add_heading("Incidencias consideradas", level=2)
    issues = document.add_table(rows=1, cols=5)
    issues.style = "Light Shading Accent 1"
    for cell, label in zip(issues.rows[0].cells, ("Clave", "Resumen", "Estado", "Responsable", "Prioridad")):
        cell.text = label
    for item in selected:
        if not isinstance(item, dict):
            continue
        cells = issues.add_row().cells
        values = (item.get("key", ""), item.get("summary", ""), item.get("status", ""), item.get("assignee") or "Sin responsable", item.get("priority", ""))
        for cell, value in zip(cells, values):
            cell.text = str(value)
    document.save(path)


def _write_pdf(path: Path, title: str, content: str, context: dict[str, object]) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=1, spaceAfter=14))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=7))
    story = [Paragraph(title, styles["ReportTitle"]), Paragraph("Análisis de calidad Jira", styles["Heading2"])]
    metrics = context.get("metrics", {})
    story.append(Paragraph(f"Incidencias en el alcance: {metrics.get('total', 0)}. Estados reales: {metrics.get('status_distribution', {})}.", styles["ReportBody"]))
    story.extend(Paragraph(_escape_pdf_text(paragraph), styles["ReportBody"]) for paragraph in _paragraphs(content))
    rows = [["Clave", "Resumen", "Estado", "Responsable"]]
    for item in context.get("selected_issues", []):
        if isinstance(item, dict):
            rows.append([str(item.get("key", "")), str(item.get("summary", ""))[:70], str(item.get("status", "")), str(item.get("assignee") or "Sin responsable")])
    if len(rows) > 1:
        table = Table(rows, colWidths=[2.2 * cm, 8.5 * cm, 3 * cm, 4 * cm], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.extend([Spacer(1, 8), table])
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.3 * cm, leftMargin=1.3 * cm, topMargin=1.3 * cm, bottomMargin=1.3 * cm).build(story)


def _paragraphs(content: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n+", content) if part.strip()]


def _escape_pdf_text(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")