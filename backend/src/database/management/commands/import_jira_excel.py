"""Load Jira historical data from an XLSX workbook without duplicate rows."""
from pathlib import Path
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.utils.dateparse import parse_datetime
from openpyxl import load_workbook
from src.database.models import JiraIssue


def value(row: dict[str, object], key: str) -> str:
    item = row.get(key)
    return "" if item is None else str(item).strip()


def jira_datetime(raw: str):
    return parse_datetime(raw.replace("Z", "+00:00")) if raw else None


class Command(BaseCommand):
    help = "Imports and updates Jira issues from DATOS_IA_HISTORICO.xlsx."

    def add_arguments(self, parser) -> None:
        parser.add_argument("--file", type=str, help="Absolute or relative XLSX path.")

    def handle(self, *args, **options) -> None:
        default_file = Path(settings.BASE_DIR).parent.parent / "DATOS_IA_HISTORICO.xlsx"
        workbook_path = Path(options["file"]) if options["file"] else default_file
        if not workbook_path.is_file():
            raise CommandError(f"No existe el archivo Excel: {workbook_path}")
        sheet = load_workbook(workbook_path, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise CommandError("El archivo Excel no contiene filas.")
        headers = [str(header).strip() if header is not None else "" for header in rows[0]]
        created = updated = skipped = 0
        for cells in rows[1:]:
            row = dict(zip(headers, cells))
            jira_id, issue_key = value(row, "ID"), value(row, "Key")
            if not jira_id or not issue_key:
                skipped += 1
                continue
            defaults = {"status": value(row, "Estado"), "status_category": value(row, "Nombre categoría estado"), "criticality": value(row, "Criticidad"), "assignee": value(row, "Nombre responsable"), "epic_key": value(row, "ID épica"), "epic_name": value(row, "Nombre épica"), "product": value(row, "Producto/Integracion"), "functionality": value(row, "Funcionalidad"), "issue_type": value(row, "Tipo incidencia"), "phase": value(row, "Fase"), "automation_status": value(row, "Estado automatización"), "description": value(row, "Descripción incidencia"), "labels": value(row, "Etiquetas"), "created_at_jira": jira_datetime(value(row, "Fecha creación")), "updated_at_jira": jira_datetime(value(row, "Fecha actualización"))}
            _, was_created = JiraIssue.objects.update_or_create(jira_id=jira_id, defaults={"issue_key": issue_key, **defaults})
            created += int(was_created)
            updated += int(not was_created)
        self.stdout.write(self.style.SUCCESS(f"Importación terminada: {created} creados, {updated} actualizados, {skipped} omitidos."))
